"""Export dishes the knowledge base does not recognize to an Excel sheet for hand labeling,
and import the filled sheet back into the menu files.

    uv run python -m michelin.parse.review_workbook export data/menus/*.json -o review.xlsx
    uv run python -m michelin.parse.review_workbook import review.xlsx data/menus/*.json

Each row is one dish; each restriction column takes 一定有 / 大概率有 / 可能有 / 没有 / 不知道
(definite / likely / possible / absent / unknown). Ingredients printed on the menu are
prefilled. Import turns the answers into components and relabels the dish with the same
rules as the knowledge base; 不知道 or a blank cell keeps that restriction unknown. Menus stay
unverified: a person still sets `verified: true`.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from michelin.parse.knowledge import PRINTED, Component, _text, label_dish, load_kb, match
from michelin.parse.label_menu import PRINTED_FIELDS
from michelin.parse.matcher import cache_key, load_cache
from michelin.schemas import Dish, Menu

COLUMNS = [  # (hit, header)
    ("pork", "猪肉 / 猪油"),
    ("beef", "牛肉"),
    ("poultry", "鸡鸭禽类"),
    ("meat", "羊肉 / 其他肉"),
    ("meat_stock", "肉汤 / 鸡粉"),
    ("fish", "鱼 / 鱼露"),
    ("shellfish", "虾蟹贝类 / 蚝油"),
    ("egg", "蛋"),
    ("dairy", "奶制品"),
    ("peanut", "花生"),
    ("tree_nut", "坚果（腰果核桃等）"),
    ("sesame", "芝麻"),
    ("soy", "大豆（豆腐酱油等）"),
    ("wheat", "小麦（面粉面条等）"),
]
ENGLISH = {  # component names shown in staff questions; UI text is English
    "pork": "pork or lard",
    "beef": "beef",
    "poultry": "chicken or duck",
    "meat": "lamb or other meat",
    "meat_stock": "meat stock or chicken powder",
    "fish": "fish or fish sauce",
    "shellfish": "shellfish or oyster sauce",
    "egg": "egg",
    "dairy": "dairy",
    "peanut": "peanuts",
    "tree_nut": "tree nuts",
    "sesame": "sesame",
    "soy": "soy",
    "wheat": "wheat",
}
ANSWERS = {"一定有": "definite", "大概率有": "likely", "可能有": "possible", "没有": None,
           "不知道": "unknown"}  # fmt: skip
FIXED = ["餐厅", "dish_id（勿改）", "英文菜名", "中文名", "价格", "菜单上的描述"]


def unrecognized(menus: list[tuple[Path, Menu]]) -> list[tuple[Menu, Dish]]:
    kb, cache = load_kb(), load_cache()
    out = []
    for _, menu in menus:
        for d in menu.dishes:
            stored = cache["matches"].get(cache_key(d))
            if match(d, kb) is None and not (stored and stored["kb_id"]):
                out.append((menu, d))
    return out


def export(menus: list[tuple[Path, Menu]], out: Path) -> int:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    fill_input = PatternFill("solid", start_color="FFF9E0")
    fill_head = PatternFill("solid", start_color="DDE8E3")
    fill_fixed = PatternFill("solid", start_color="F2F2F2")
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    guide = wb.active
    guide.title = "填写说明"
    rows = [
        ["菜品标注表：知识库认不出来的菜"],
        [],
        ["怎么填"],
        ["1. 只填「菜品」页里浅黄色的格子：每道菜 × 每类成分，点格子选下拉选项。"],
        ["2. 一定有：这道菜的定义就有，或菜单上写明了。大概率有：多数做法都有。"],
        ["   可能有：有些做法会放（如蚝油、鸡粉、猪油）。没有：确定没有。不知道：拿不准。"],
        ["3. 拿不准就选「不知道」，不要猜。不知道 = 系统会让忌口的人先问服务员。"],
        ["4. 已经预填了「一定有」的格子，是菜单上印出来的成分，可以改。"],
        ["5. 灰色列不要改。最后一列「备注」可以写问题或理由。"],
        ["6. 填完把文件发回来，我导回菜单文件。导回后菜单仍需你确认 verified: true。"],
        [],
        ["示例（一行填好的样子）"],
        [
            "菜",
            "猪肉 / 猪油",
            "牛肉",
            "鸡鸭禽类",
            "肉汤 / 鸡粉",
            "虾蟹贝类 / 蚝油",
            "蛋",
            "花生",
            "芝麻",
            "大豆",
            "小麦",
        ],
        [
            "酸辣汤 Hot and Sour Soup",
            "大概率有",
            "没有",
            "没有",
            "大概率有",
            "没有",
            "大概率有",
            "没有",
            "可能有",
            "一定有",
            "大概率有",
        ],
    ]
    for r in rows:
        guide.append(r)
    for row in guide.iter_rows():
        for cell in row:
            cell.font = font
    for ref in ("A1", "A3", "A12"):
        guide[ref].font = Font(name="Arial", size=12 if ref == "A1" else 10, bold=True)
    for cell in guide[13]:
        cell.font, cell.fill = bold, fill_head
    for cell in guide[14][1:]:
        cell.fill = fill_input
    guide.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJK":
        guide.column_dimensions[col].width = 13

    ws = wb.create_sheet("菜品")
    headers = FIXED + [h for _, h in COLUMNS] + ["备注"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = bold, fill_head, wrap
    dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "请从下拉选项里选：" + " / ".join(ANSWERS), "选项不对"
    ws.add_data_validation(dv)

    dishes = unrecognized(menus)
    for i, (menu, d) in enumerate(dishes, start=2):
        text = _text(d)  # misleading names ("fish-fragrant", 鱼香) removed first
        printed = {h for h, pat in PRINTED.items() if re.search(pat, text, re.IGNORECASE)}
        ws.append([menu.restaurant_name, f"{menu.restaurant_id}/{d.id}", d.name_en, d.name_zh,
                   d.price, d.description_raw]
                  + ["一定有" if h in printed else None for h, _ in COLUMNS] + [None])  # fmt: skip
        for j, cell in enumerate(ws[i], start=1):
            cell.font, cell.alignment = font, wrap
            if j <= len(FIXED):
                cell.fill = fill_fixed
            elif j <= len(FIXED) + len(COLUMNS):
                cell.fill = fill_input
                if cell.value:
                    cell.comment = Comment("菜单上写明，已预填", "Michelin")
        ws.cell(i, 5).number_format = "$0.00"
    first, last = len(FIXED) + 1, len(FIXED) + len(COLUMNS)
    from openpyxl.utils import get_column_letter as col

    dv.add(f"{col(first)}2:{col(last)}{len(dishes) + 1}")
    widths = [12, 16, 28, 14, 8, 30] + [9] * len(COLUMNS) + [30]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[col(j)].width = w
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "G2"
    ws.column_dimensions["B"].hidden = False
    wb.active = 1
    wb.save(out)
    return len(dishes)


def components_from_row(row: dict) -> tuple[list[Component], bool]:
    """(components, fully answered). Unknown and blank answers stay out of the components."""
    comps, complete = [], True
    for hit, header in COLUMNS:
        answer = (row.get(header) or "").strip()
        cert = ANSWERS.get(answer, "unknown") if answer else "unknown"
        if cert == "unknown":
            complete = False
        elif cert:
            comps.append(Component(ENGLISH[hit], (hit,), cert))
    return comps, complete


def import_sheet(sheet: Path, menu_paths: list[Path]) -> list[str]:
    from openpyxl import load_workbook

    ws = load_workbook(sheet)["菜品"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    answers = {r[1]: dict(zip(header, r, strict=False)) for r in rows[1:] if r[1]}
    log = []
    for path in menu_paths:
        menu = Menu.model_validate_json(path.read_text(encoding="utf-8"))
        dishes = []
        for d in menu.dishes:
            row = answers.get(f"{menu.restaurant_id}/{d.id}")
            if row is None:
                dishes.append(d)
                continue
            comps, complete = components_from_row(row)
            entry = {"id": f"reviewed:{d.id}", "name_zh": [], "aliases": [],
                     "category": d.category.value, "portion": d.portion.value,
                     "components": [{"name": c.name, "hits": list(c.hits), "certainty": c.certainty}
                                    for c in comps]}  # fmt: skip
            # Start from what the menu prints, so earlier labels and questions do not linger.
            printed = Dish(**{k: getattr(d, k) for k in PRINTED_FIELDS})
            labeled = label_dish(printed, reviewed=entry if complete else None)
            dish = labeled.dish
            if not complete:  # partial answers: add them on top of the unknown baseline
                dish = _apply_partial(dish, comps)
            note = (row.get("备注") or "").strip()
            if note:
                dish = dish.model_copy(
                    update={
                        "confirm_with_staff": [*dish.confirm_with_staff, f"{d.name_en}: {note}"]
                    }
                )
            dishes.append(dish)
            log.append(f"{menu.restaurant_id}/{d.id}: {'complete' if complete else 'partial'}")
        menu = menu.model_copy(update={"dishes": dishes, "verified": False})
        path.write_text(menu.model_dump_json(indent=2) + "\n", encoding="utf-8")
    return log


def _apply_partial(dish: Dish, comps: list[Component]) -> Dish:
    """Known-present answers become inferred flags and block diets; everything else stays
    unknown, so nothing is certified from a half-filled row."""
    from michelin.parse.knowledge import VEGAN_BREAKERS, VEGETARIAN_BREAKERS
    from michelin.schemas import Allergen, AllergenFlag, EvidenceTier

    flags = {f.allergen: f for f in dish.allergens}
    severe = [c for c in comps if c.certainty in ("definite", "likely")]
    for c in severe:
        for hit in c.hits:
            if hit in Allergen._value2member_map_:
                flags[Allergen(hit)] = AllergenFlag(allergen=Allergen(hit), tier=EvidenceTier.INFERRED,
                                                    confidence=0.8, reason=f"{c.name}")  # fmt: skip
    hits = {h for c in severe for h in c.hits}
    update = {"allergens": sorted(flags.values(), key=lambda f: f.allergen.value)}
    if hits & VEGETARIAN_BREAKERS:
        update["is_vegetarian"] = False
    if hits & VEGAN_BREAKERS:
        update["is_vegan"] = False
    if "pork" in hits:
        update["contains_pork"] = True
    if "beef" in hits:
        update["contains_beef"] = True
    return dish.model_copy(update=update)


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("menus", nargs="+", type=Path)
    e.add_argument("-o", "--out", type=Path, required=True)
    i = sub.add_parser("import")
    i.add_argument("sheet", type=Path)
    i.add_argument("menus", nargs="+", type=Path)
    args = ap.parse_args(argv)
    if args.cmd == "export":
        menus = [(p, Menu.model_validate_json(p.read_text(encoding="utf-8"))) for p in args.menus]
        print(f"Wrote {export(menus, args.out)} dishes to {args.out}")
    else:
        print("\n".join(import_sheet(args.sheet, args.menus)) or "no rows matched")


if __name__ == "__main__":
    main()
