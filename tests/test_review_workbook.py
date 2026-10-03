from openpyxl import load_workbook

from michelin.parse.review_workbook import COLUMNS, export, import_sheet
from michelin.schemas import Allergen, EvidenceTier, Menu


def fill(path, dish_key, answers, note=None):
    wb = load_workbook(path)
    ws = wb["菜品"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[1].value == dish_key:
            for header_name, value in answers.items():
                row[header.index(header_name)].value = value
            if note:
                row[header.index("备注")].value = note
    wb.save(path)


def test_roundtrip_complete_row_certifies_and_partial_row_does_not(tmp_path, sample_menu):
    menu_path = tmp_path / "menu.json"
    menu_path.write_text(sample_menu.model_dump_json(), encoding="utf-8")
    sheet = tmp_path / "review.xlsx"
    assert export([(menu_path, sample_menu)], sheet) >= 1  # dry_pot_cauliflower is unknown

    key = f"{sample_menu.restaurant_id}/dry_pot_cauliflower"
    answers = {h: "没有" for _, h in COLUMNS}
    answers |= {
        "猪肉 / 猪油": "大概率有",
        "大豆（豆腐酱油等）": "一定有",
        "小麦（面粉面条等）": "可能有",
    }
    fill(sheet, key, answers, note="Ask if pork belly can be left out")
    assert import_sheet(sheet, [menu_path]) == [f"{key}: complete"]

    dish = Menu.model_validate_json(menu_path.read_text(encoding="utf-8")).dish(
        "dry_pot_cauliflower"
    )
    assert dish.is_vegetarian is False and dish.contains_pork is True
    tiers = {f.allergen: f.tier for f in dish.allergens}
    assert tiers[Allergen.SOY] == EvidenceTier.INFERRED
    assert tiers[Allergen.WHEAT] == EvidenceTier.UNKNOWN
    assert Allergen.SHELLFISH not in tiers  # answered 没有
    assert any("pork belly" in q for q in dish.confirm_with_staff)


def test_partial_row_only_adds_what_is_known(tmp_path, sample_menu):
    menu_path = tmp_path / "menu.json"
    menu_path.write_text(sample_menu.model_dump_json(), encoding="utf-8")
    sheet = tmp_path / "review.xlsx"
    export([(menu_path, sample_menu)], sheet)
    key = f"{sample_menu.restaurant_id}/dry_pot_cauliflower"
    fill(sheet, key, {"猪肉 / 猪油": "一定有"})
    assert import_sheet(sheet, [menu_path]) == [f"{key}: partial"]
    dish = Menu.model_validate_json(menu_path.read_text(encoding="utf-8")).dish(
        "dry_pot_cauliflower"
    )
    assert dish.contains_pork is True
    assert dish.contains_beef is None  # unanswered stays unknown
    assert {f.tier for f in dish.allergens if f.allergen == Allergen.SHELLFISH} == {
        EvidenceTier.UNKNOWN
    }


def test_reviewed_row_overrides_a_knowledge_base_match(tmp_path, sample_menu):
    from openpyxl import Workbook

    from michelin.parse.review_workbook import FIXED

    menu_path = tmp_path / "menu.json"
    menu_path.write_text(sample_menu.model_dump_json(), encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "菜品"
    ws.append(FIXED + [h for _, h in COLUMNS] + ["备注"])
    answers = {h: "没有" for _, h in COLUMNS} | {"羊肉 / 其他肉": "可能有", "猪肉 / 猪油": "可能有",
                                              "大豆（豆腐酱油等）": "一定有"}  # fmt: skip
    ws.append([None, f"{sample_menu.restaurant_id}/mapo_tofu", None, None, None, None]
              + [answers[h] for _, h in COLUMNS] + ["Menu marks it vegetarian"])  # fmt: skip
    sheet = tmp_path / "answers.xlsx"
    wb.save(sheet)
    import_sheet(sheet, [menu_path])
    dish = Menu.model_validate_json(menu_path.read_text(encoding="utf-8")).dish("mapo_tofu")
    assert dish.is_vegetarian is True  # possible meat: vegetarian may ask
    assert dish.contains_pork is False  # diets ask: no-pork may order after the question
    assert any("pork" in q for q in dish.confirm_with_staff)
    assert not any("same dish" in q or "recognize" in q for q in dish.confirm_with_staff)
